"""
Presentation Deck Generator
Bangkok Airbnb Market Intelligence
"""
import sys
sys.path.insert(0, '.')

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import io
import os

os.makedirs("presentation", exist_ok=True)

# Colors
DARK_BLUE = RGBColor(0x1F, 0x38, 0x64)
MED_BLUE = RGBColor(0x2E, 0x75, 0xB6)
CYAN = RGBColor(0x00, 0xD4, 0xFF)
GREEN = RGBColor(0x4A, 0xDE, 0x80)
RED = RGBColor(0xFF, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF8, 0xF9, 0xFA)
GRAY = RGBColor(0x7A, 0x84, 0x99)
ORANGE = RGBColor(0xF5, 0x9E, 0x0B)

# Load data
df = pd.read_csv("data/processed/listings_enriched.csv", low_memory=False)
df["occupancy_rate"] = 1 - (df["availability_365"] / 365)
df["estimated_annual_revenue"] = df["price"] * (365 - df["availability_365"])
df_clean = df[df["price"].between(100, 10000)].copy()
reviews = pd.read_csv("data/processed/reviews_clean.csv", low_memory=False)
reviews["date"] = pd.to_datetime(reviews["date"])

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

blank = prs.slide_layouts[6]  # blank layout

def add_rect(slide, l, t, w, h, fill, alpha=None):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    return shape

def add_text(slide, text, l, t, w, h, size=18, bold=False,
             color=WHITE, align=PP_ALIGN.LEFT, wrap=True):
    txb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txb

def add_chart_image(slide, fig, l, t, w, h):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", facecolor="white")
    buf.seek(0)
    slide.shapes.add_picture(buf, Inches(l), Inches(t), Inches(w), Inches(h))
    plt.close(fig)

# ============ SLIDE 1: COVER ============
slide = prs.slides.add_slide(blank)

# Dark background
add_rect(slide, 0, 0, 13.33, 7.5, DARK_BLUE)
# Cyan accent bar
add_rect(slide, 0, 5.8, 13.33, 0.15, CYAN)

add_text(slide, "BANGKOK AIRBNB", 1, 1.5, 11, 1.2, size=52,
         bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, "Market Intelligence Report", 1, 2.8, 11, 0.8, size=28,
         color=CYAN, align=PP_ALIGN.CENTER)
add_text(slide, "Expernetic Data Engineer Intern Assessment  |  Primesh Marasinghe  |  June 2026",
         1, 3.7, 11, 0.5, size=14, color=GRAY, align=PP_ALIGN.CENTER)

# KPI boxes
kpis = [("28,806", "Listings"), ("฿1,359", "Median Price"),
        ("31.5%", "Avg Occupancy"), ("583K", "Reviews")]
for i, (val, label) in enumerate(kpis):
    x = 1.2 + i * 2.8
    add_rect(slide, x, 4.4, 2.4, 1.2, MED_BLUE)
    add_text(slide, val, x, 4.5, 2.4, 0.6, size=24, bold=True,
             color=CYAN, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, 5.1, 2.4, 0.4, size=11,
             color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

add_text(slide, "Inside Airbnb Dataset  ·  Bangkok, Thailand  ·  September 2025 Scrape",
         1, 6.1, 11, 0.4, size=11, color=GRAY, align=PP_ALIGN.CENTER)

# ============ SLIDE 2: AGENDA ============
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, 13.33, 7.5, RGBColor(0x0e, 0x11, 0x17))
add_rect(slide, 0, 0, 13.33, 1.2, DARK_BLUE)
add_rect(slide, 0, 1.2, 0.08, 6.3, CYAN)

add_text(slide, "Presentation Agenda", 0.3, 0.2, 12, 0.8,
         size=28, bold=True, color=WHITE)

agenda = [
    ("01", "Market Overview", "Scale, structure, and key market characteristics"),
    ("02", "Pricing Analysis", "Price drivers, room type premiums, neighbourhood dynamics"),
    ("03", "Host Intelligence", "Market concentration, superhost impact, segmentation"),
    ("04", "Demand & Seasonality", "Peak seasons, COVID recovery, booking patterns"),
    ("05", "Statistical Findings", "5 hypothesis tests with effect sizes"),
    ("06", "ML & AI", "Price prediction, NLP, RAG system, recommendations"),
    ("07", "Business Recommendations", "Actionable insights for hosts, investors, operators"),
]

for i, (num, title, desc) in enumerate(agenda):
    y = 1.4 + i * 0.78
    add_rect(slide, 0.3, y, 0.55, 0.55, MED_BLUE)
    add_text(slide, num, 0.3, y, 0.55, 0.55, size=14,
             bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, title, 1.1, y, 4, 0.3, size=14,
             bold=True, color=CYAN)
    add_text(slide, desc, 1.1, y + 0.28, 11, 0.3, size=10, color=GRAY)

# ============ SLIDE 3: MARKET OVERVIEW ============
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, 13.33, 7.5, RGBColor(0x0e, 0x11, 0x17))
add_rect(slide, 0, 0, 13.33, 1.2, DARK_BLUE)
add_rect(slide, 0, 1.2, 0.08, 6.3, CYAN)

add_text(slide, "01  Market Overview", 0.3, 0.2, 12, 0.8,
         size=28, bold=True, color=WHITE)

# Market stats
stats = [
    ("28,806", "Active Listings", MED_BLUE),
    ("8,874", "Unique Hosts", MED_BLUE),
    ("50", "Neighbourhoods", MED_BLUE),
    ("10.5M", "Calendar Records", MED_BLUE),
]
for i, (val, label, col) in enumerate(stats):
    x = 0.4 + i * 3.2
    add_rect(slide, x, 1.4, 2.9, 1.5, col)
    add_text(slide, val, x, 1.5, 2.9, 0.8, size=30,
             bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, 2.3, 2.9, 0.4, size=11,
             color=WHITE, align=PP_ALIGN.CENTER)

# Key findings
findings = [
    "🏙️  Bangkok is one of SE Asia's most dynamic short-term rental markets",
    "📈  2024 review volume hit all-time high: 148,869 reviews (+152% vs 2019 pre-COVID)",
    "🏢  Market dominated by commercial operators: 10.6% of hosts control 57% of listings",
    "💰  Median nightly price: ฿1,359 THB (~$38 USD) — affordable by global standards",
    "⭐  81.4% of listings rated above 4.5 — classic rating inflation pattern confirmed",
    "🚇  BTS Skytrain proximity is the #1 guest decision factor (2,513 review mentions)",
]
for i, finding in enumerate(findings):
    y = 3.1 + i * 0.62
    add_text(slide, finding, 0.4, y, 12.5, 0.5, size=12, color=WHITE)

# ============ SLIDE 4: PRICING ANALYSIS ============
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, 13.33, 7.5, RGBColor(0x0e, 0x11, 0x17))
add_rect(slide, 0, 0, 13.33, 1.2, DARK_BLUE)
add_rect(slide, 0, 1.2, 0.08, 6.3, CYAN)
add_text(slide, "02  Pricing Analysis", 0.3, 0.2, 12, 0.8,
         size=28, bold=True, color=WHITE)

fig, axes = plt.subplots(1, 2, figsize=(10, 4))
fig.patch.set_facecolor("#1a1f2e")
for ax in axes:
    ax.set_facecolor("#1a1f2e")
    ax.tick_params(colors="white")
    ax.spines["bottom"].set_color("#2a2f3e")
    ax.spines["left"].set_color("#2a2f3e")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)

room_price = df_clean.groupby("room_type")["price"].median().sort_values()
bars = axes[0].barh(room_price.index, room_price.values,
                    color=["#2E75B6","#4ade80","#f59e0b","#ff6b6b"])
axes[0].set_title("Median Price by Room Type (THB)", color="white", fontsize=11)
axes[0].set_xlabel("Median Price (฿)", color="white")
for bar, val in zip(bars, room_price.values):
    axes[0].text(bar.get_width()+20, bar.get_y()+bar.get_height()/2,
                 f"฿{val:,.0f}", va="center", fontsize=9, color="white")

top_nb = df_clean.groupby("neighbourhood_cleansed")["price"].median().sort_values(ascending=False).head(8)
axes[1].barh(top_nb.index[::-1], top_nb.values[::-1], color="#2E75B6")
axes[1].set_title("Top 8 Neighbourhoods by Median Price", color="white", fontsize=11)
axes[1].set_xlabel("Median Price (฿)", color="white")
plt.tight_layout()
add_chart_image(slide, fig, 0.3, 1.4, 8.5, 4.5)

insights = [
    "38% premium",
    "for entire homes",
    "vs private rooms",
    "(p<0.0001)",
    "",
    "Parthum Wan",
    "฿2,225 highest",
    "",
    "TV amenity adds",
    "+68.6% premium",
]
add_rect(slide, 9.1, 1.4, 3.9, 4.5, RGBColor(0x1a, 0x1f, 0x2e))
add_text(slide, "💡 Key Insights", 9.2, 1.5, 3.7, 0.4,
         size=12, bold=True, color=CYAN)
insight_texts = [
    "• Entire homes: 38% premium over private rooms (p<0.0001)",
    "• Parthum Wan commands highest prices at ฿2,225 median",
    "• TV amenity adds +68.6% price premium",
    "• Hot tub: +49.8% | Kitchen: +36.9%",
    "• Neighbourhood explains 10.5% of price variance (η²=0.11)",
]
for i, txt in enumerate(insight_texts):
    add_text(slide, txt, 9.2, 2.0 + i*0.65, 3.7, 0.55, size=10, color=WHITE)

# ============ SLIDE 5: HOST INTELLIGENCE ============
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, 13.33, 7.5, RGBColor(0x0e, 0x11, 0x17))
add_rect(slide, 0, 0, 13.33, 1.2, DARK_BLUE)
add_rect(slide, 0, 1.2, 0.08, 6.3, CYAN)
add_text(slide, "03  Host Market Intelligence", 0.3, 0.2, 12, 0.8,
         size=28, bold=True, color=WHITE)

# Superhost comparison boxes
add_rect(slide, 0.4, 1.4, 5.8, 2.5, RGBColor(0x1a, 0x1f, 0x2e))
add_rect(slide, 6.5, 1.4, 6.5, 2.5, RGBColor(0x1a, 0x2e, 0x1a))

add_text(slide, "Non-Superhost", 0.5, 1.5, 5.5, 0.4, size=13, color=GRAY, bold=True)
add_text(slide, "฿25,968", 0.5, 1.9, 5.5, 0.7, size=36, bold=True, color=RED)
add_text(slide, "Median Annual Revenue", 0.5, 2.6, 5.5, 0.3, size=10, color=GRAY)
add_text(slide, "Avg Rating: 4.575", 0.5, 3.0, 5.5, 0.3, size=11, color=WHITE)

add_text(slide, "⭐ Superhost", 6.6, 1.5, 6, 0.4, size=13, color=CYAN, bold=True)
add_text(slide, "฿114,431", 6.6, 1.9, 6, 0.7, size=36, bold=True, color=GREEN)
add_text(slide, "Median Annual Revenue  (+340%)", 6.6, 2.6, 6, 0.3, size=10, color=GRAY)
add_text(slide, "Avg Rating: 4.858  (Cohen's d=0.54)", 6.6, 3.0, 6, 0.3, size=11, color=WHITE)

# Market concentration
add_rect(slide, 0.4, 4.1, 12.5, 2.9, RGBColor(0x1a, 0x1f, 0x2e))
add_text(slide, "⚠️  Market Concentration Alert", 0.6, 4.2, 12, 0.4,
         size=13, bold=True, color=ORANGE)

conc_stats = [
    ("10.6%\nof hosts", "Control\n57% of listings"),
    ("Top host\n243 listings", "Single operator"),
    ("6 host\nsegments", "K-Means clustering"),
    ("62.7%\ncasual hosts", "Single listing only"),
]
for i, (val, label) in enumerate(conc_stats):
    x = 0.6 + i * 3.1
    add_text(slide, val, x, 4.7, 2.8, 0.7, size=18, bold=True,
             color=CYAN, align=PP_ALIGN.CENTER)
    add_text(slide, label, x, 5.4, 2.8, 0.4, size=10,
             color=GRAY, align=PP_ALIGN.CENTER)

# ============ SLIDE 6: DEMAND & SEASONALITY ============
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, 13.33, 7.5, RGBColor(0x0e, 0x11, 0x17))
add_rect(slide, 0, 0, 13.33, 1.2, DARK_BLUE)
add_rect(slide, 0, 1.2, 0.08, 6.3, CYAN)
add_text(slide, "04  Demand & Seasonality", 0.3, 0.2, 12, 0.8,
         size=28, bold=True, color=WHITE)

fig, axes = plt.subplots(1, 2, figsize=(11, 4))
fig.patch.set_facecolor("#1a1f2e")
for ax in axes:
    ax.set_facecolor("#1a1f2e")
    ax.tick_params(colors="white", labelsize=8)
    for spine in ax.spines.values():
        spine.set_color("#2a2f3e")

monthly = reviews[reviews["date"].dt.year.between(2019,2025)].groupby(
    reviews["date"].dt.to_period("M")).size().reset_index()
monthly.columns = ["period","count"]
monthly["date"] = monthly["period"].dt.to_timestamp()
axes[0].fill_between(monthly["date"], monthly["count"], alpha=0.3, color="#00d4ff")
axes[0].plot(monthly["date"], monthly["count"], color="#00d4ff", linewidth=2)
axes[0].axvline(pd.Timestamp("2020-03-01"), color="#ff6b6b", linestyle="--", alpha=0.8)
axes[0].text(pd.Timestamp("2020-06-01"), monthly["count"].max()*0.7,
             "COVID-19", color="#ff6b6b", fontsize=8)
axes[0].set_title("Monthly Review Volume 2019–2025", color="white", fontsize=11)
axes[0].set_ylabel("Monthly Reviews", color="white")

month_names = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
occ_by_month = [27.7,24.4,23.4,28.3,27.6,30.5,41.8,41.6,43.2,33.4,27.7,27.6]
colors_m = ["#ff6b6b" if o > 40 else "#2E75B6" for o in occ_by_month]
axes[1].bar(month_names, occ_by_month, color=colors_m)
axes[1].set_title("Monthly Occupancy Rate (%)", color="white", fontsize=11)
axes[1].set_ylabel("Occupancy Rate (%)", color="white")
axes[1].tick_params(axis="x", rotation=45)
plt.tight_layout()
add_chart_image(slide, fig, 0.3, 1.4, 9.5, 4.5)

add_rect(slide, 10.1, 1.4, 3.0, 4.5, RGBColor(0x1a, 0x1f, 0x2e))
add_text(slide, "📊 Key Demand\nInsights", 10.2, 1.5, 2.8, 0.7,
         size=12, bold=True, color=CYAN)
demand_insights = [
    "Peak: Jul–Sep\n(43% occupancy)",
    "Trough: Feb–Mar\n(~24%)",
    "2024: All-time\nhigh 148,869",
    "+152% vs 2019\npre-COVID",
    "No weekend\npremium (p=0.526)",
]
for i, txt in enumerate(demand_insights):
    add_text(slide, f"• {txt}", 10.2, 2.3 + i*0.68, 2.8, 0.6, size=10, color=WHITE)

# ============ SLIDE 7: STATISTICAL FINDINGS ============
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, 13.33, 7.5, RGBColor(0x0e, 0x11, 0x17))
add_rect(slide, 0, 0, 13.33, 1.2, DARK_BLUE)
add_rect(slide, 0, 1.2, 0.08, 6.3, CYAN)
add_text(slide, "05  Statistical Findings", 0.3, 0.2, 12, 0.8,
         size=28, bold=True, color=WHITE)

hypotheses = [
    ("H1", "Entire home > Private room price", "REJECT H₀", "p<0.0001", "r=0.29 moderate", GREEN),
    ("H2", "Superhost > Non-superhost rating", "REJECT H₀", "p<0.0001", "d=0.54 medium", GREEN),
    ("H3", ">10 reviews affects price", "REJECT H₀*", "p=0.000005", "r=0.037 negligible", ORANGE),
    ("H4", "Neighbourhood price differences", "REJECT H₀", "p<0.0001", "η²=0.11 moderate", GREEN),
    ("H5", "Weekend vs weekday pricing", "FAIL TO REJECT", "p=0.526", "d=0.012 negligible", RED),
]

headers = ["ID", "Hypothesis", "Decision", "P-value", "Effect Size"]
col_widths = [0.7, 5.2, 2.2, 1.8, 2.3]
col_starts = [0.3]
for w in col_widths[:-1]:
    col_starts.append(col_starts[-1] + w)

# Header row
add_rect(slide, 0.3, 1.4, 12.5, 0.45, DARK_BLUE)
for j, (h, x) in enumerate(zip(headers, col_starts)):
    add_text(slide, h, x+0.05, 1.42, col_widths[j]-0.1, 0.4,
             size=11, bold=True, color=CYAN)

for i, (hid, hyp, decision, pval, effect, col) in enumerate(hypotheses):
    y = 1.9 + i * 0.85
    bg = RGBColor(0x1a, 0x1f, 0x2e) if i % 2 == 0 else RGBColor(0x14, 0x19, 0x28)
    add_rect(slide, 0.3, y, 12.5, 0.8, bg)
    row = [hid, hyp, decision, pval, effect]
    for j, (val, x) in enumerate(zip(row, col_starts)):
        c = col if j == 2 else WHITE
        bold = j == 2
        add_text(slide, val, x+0.05, y+0.1, col_widths[j]-0.1, 0.6,
                 size=10, bold=bold, color=c)

add_text(slide, "* H3: Statistically significant but effect size negligible — large n amplifies trivial differences",
         0.3, 6.25, 12.5, 0.35, size=9, color=GRAY)

# ============ SLIDE 8: ML & AI ============
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, 13.33, 7.5, RGBColor(0x0e, 0x11, 0x17))
add_rect(slide, 0, 0, 13.33, 1.2, DARK_BLUE)
add_rect(slide, 0, 1.2, 0.08, 6.3, CYAN)
add_text(slide, "06  Machine Learning & AI", 0.3, 0.2, 12, 0.8,
         size=28, bold=True, color=WHITE)

# ML models
add_text(slide, "Price Prediction (3 Models + 5-Fold CV)", 0.3, 1.35, 7, 0.4,
         size=13, bold=True, color=CYAN)
models_data = [
    ("Ridge Regression", "฿655", "0.449", "42.1%", RED),
    ("Random Forest", "฿501", "0.652", "29.2%", GREEN),
    ("Gradient Boosting ✓", "฿504", "0.662", "29.1%", CYAN),
]
headers_ml = ["Model", "MAE", "R²", "MAPE"]
col_w = [3.2, 1.3, 1.2, 1.3]
col_x = [0.3]
for w in col_w[:-1]:
    col_x.append(col_x[-1]+w)

add_rect(slide, 0.3, 1.8, 7.2, 0.4, DARK_BLUE)
for j, (h, x, w) in enumerate(zip(headers_ml, col_x, col_w)):
    add_text(slide, h, x+0.05, 1.82, w, 0.35, size=10, bold=True, color=WHITE)

for i, (name, mae, r2, mape, col) in enumerate(models_data):
    y = 2.25 + i*0.6
    bg = RGBColor(0x1a,0x1f,0x2e) if i%2==0 else RGBColor(0x14,0x19,0x28)
    add_rect(slide, 0.3, y, 7.2, 0.55, bg)
    for j, (val, x, w) in enumerate(zip([name,mae,r2,mape], col_x, col_w)):
        c = col if j==0 else WHITE
        add_text(slide, val, x+0.05, y+0.08, w, 0.4, size=10, color=c)

# NLP & AI
add_text(slide, "NLP & AI Components", 8.0, 1.35, 5, 0.4, size=13, bold=True, color=CYAN)
ai_items = [
    ("🎭", "VADER Sentiment", "72.9% positive · r=0.089 vs scores"),
    ("📚", "LDA Topic Modeling", "6 topics · 40% non-English detected"),
    ("🔍", "Named Entity Recognition", "BTS #1 landmark · Small #1 complaint"),
    ("🤖", "RAG Q&A System", "TF-IDF retrieval over 10K reviews"),
    ("🎯", "Recommendation System", "Cosine similarity · 0.99+ scores"),
    ("💡", "Listing Advisor", "Data-driven improvement recommendations"),
]
for i, (icon, title, desc) in enumerate(ai_items):
    y = 1.8 + i*0.75
    add_rect(slide, 8.0, y, 5.0, 0.65, RGBColor(0x1a,0x1f,0x2e))
    add_text(slide, f"{icon} {title}", 8.1, y+0.05, 4.8, 0.3, size=11, bold=True, color=CYAN)
    add_text(slide, desc, 8.1, y+0.33, 4.8, 0.25, size=9, color=GRAY)

# Feature importance
add_text(slide, "Top Price Drivers (Gradient Boosting):", 0.3, 4.4, 7, 0.35,
         size=11, bold=True, color=WHITE)
features = [("Bedrooms", 43.2), ("Neighbourhood Median Price", 16.9),
            ("Accommodates", 12.2), ("Host Tenure", 4.3), ("Host Listings", 4.3)]
for i, (feat, imp) in enumerate(features):
    y = 4.8 + i*0.42
    add_rect(slide, 0.3, y, imp/10, 0.32, MED_BLUE)
    add_text(slide, f"{feat}: {imp}%", 0.4, y+0.04, 6.5, 0.28, size=9, color=WHITE)

# ============ SLIDE 9: RECOMMENDATIONS ============
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, 13.33, 7.5, RGBColor(0x0e, 0x11, 0x17))
add_rect(slide, 0, 0, 13.33, 1.2, DARK_BLUE)
add_rect(slide, 0, 1.2, 0.08, 6.3, CYAN)
add_text(slide, "07  Business Recommendations", 0.3, 0.2, 12, 0.8,
         size=28, bold=True, color=WHITE)

sections = [
    ("🏠 FOR HOSTS", MED_BLUE, [
        "Pursue Superhost status — revenue jumps 340% (฿25,968 → ฿114,431)",
        "Implement weekend pricing premium — zero competition currently",
        "Highlight BTS/MRT proximity — #1 guest decision factor",
        "Upgrade amenities: TV (+68.6%), Hot tub (+49.8%), Kitchen (+36.9%)",
        "Dynamic pricing for Jul–Sep peak (43% vs 24% Feb trough)",
    ]),
    ("📈 FOR INVESTORS", GREEN, [
        "Bangkok market at all-time high — 2024 reviews +152% vs 2019",
        "Target Bang Rak / Vadhana for premium positioning",
        "Mid-market ฿1,000–2,500 captures peak demand concentration",
        "High Demand segment achieves 61% occupancy",
        "Avoid shared rooms — only 11.6% occupancy rate",
    ]),
    ("🏢 FOR OPERATORS", ORANGE, [
        "Monitor commercial concentration — 57% supply risk",
        "Address rating inflation — 81.4% above 4.5",
        "Deploy multilingual NLP for full review coverage",
        "Introduce weekend pricing tools for hosts",
        "Proactive regulatory engagement — Bangkok mirrors Amsterdam",
    ]),
]

for i, (title, col, recs) in enumerate(sections):
    x = 0.3 + i * 4.35
    add_rect(slide, x, 1.4, 4.1, 5.7, RGBColor(0x1a, 0x1f, 0x2e))
    add_rect(slide, x, 1.4, 4.1, 0.5, DARK_BLUE)
    add_text(slide, title, x+0.1, 1.45, 3.9, 0.4, size=12, bold=True, color=col)
    for j, rec in enumerate(recs):
        add_text(slide, f"• {rec}", x+0.1, 2.05+j*0.95, 3.9, 0.85,
                 size=9, color=WHITE)

# ============ SLIDE 10: TECH STACK ============
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, 13.33, 7.5, RGBColor(0x0e, 0x11, 0x17))
add_rect(slide, 0, 0, 13.33, 1.2, DARK_BLUE)
add_rect(slide, 0, 1.2, 0.08, 6.3, CYAN)
add_text(slide, "Technical Architecture & Stack", 0.3, 0.2, 12, 0.8,
         size=28, bold=True, color=WHITE)

tech_sections = [
    ("⚙️ Pipeline", ["Python 3.11", "pandas / numpy", "DuckDB SQL", "dbt models", "Docker"]),
    ("📊 Analysis", ["scipy / statsmodels", "scikit-learn", "matplotlib / seaborn", "plotly", "VADER NLP"]),
    ("🤖 AI/ML", ["Gradient Boosting", "K-Means Clustering", "TF-IDF RAG", "LDA Topics", "Rule-based NER"]),
    ("🎨 Products", ["Streamlit Dashboard", "AI Analyst Chatbot", "ReportLab PDF", "Stream Sim", "PPTX Deck"]),
    ("☁️ Architecture", ["AWS Cloud Design", "Kafka Streaming", "Star Schema", "21 Unit Tests", "Git + GitHub"]),
]

for i, (title, items) in enumerate(tech_sections):
    x = 0.3 + i * 2.6
    add_rect(slide, x, 1.4, 2.4, 5.6, RGBColor(0x1a, 0x1f, 0x2e))
    add_rect(slide, x, 1.4, 2.4, 0.5, DARK_BLUE)
    add_text(slide, title, x+0.1, 1.45, 2.2, 0.4, size=11, bold=True, color=CYAN)
    for j, item in enumerate(items):
        add_text(slide, f"• {item}", x+0.1, 2.05+j*0.9, 2.2, 0.8,
                 size=10, color=WHITE)

add_text(slide, "GitHub: github.com/primemarasinghe/airbnb-market-intelligence",
         0.3, 7.1, 12.5, 0.3, size=10, color=GRAY, align=PP_ALIGN.CENTER)

# ============ SLIDE 11: CLOSING ============
slide = prs.slides.add_slide(blank)
add_rect(slide, 0, 0, 13.33, 7.5, DARK_BLUE)
add_rect(slide, 0, 6.8, 13.33, 0.7, RGBColor(0x00, 0xD4, 0xFF))

add_text(slide, "Thank You", 1, 1.5, 11, 1.5, size=52,
         bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, "Bangkok Airbnb Market Intelligence", 1, 3.0, 11, 0.7,
         size=20, color=CYAN, align=PP_ALIGN.CENTER)

summary = [
    "✅ 8 Jupyter Notebooks  |  ✅ 21 Unit Tests  |  ✅ Docker Containerized",
    "✅ Streamlit Dashboard  |  ✅ AI Analyst Chatbot  |  ✅ Automated PDF Reports",
    "✅ dbt Models  |  ✅ Stream Processing Sim  |  ✅ AWS Architecture Design",
]
for i, txt in enumerate(summary):
    add_text(slide, txt, 1, 3.8+i*0.6, 11, 0.5, size=12,
             color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

add_text(slide, "Primesh Marasinghe  |  BSc (Hons) Data Science, SLIIT  |  June 2026",
         1, 5.5, 11, 0.4, size=13, color=GRAY, align=PP_ALIGN.CENTER)
add_text(slide, "github.com/primemarasinghe/airbnb-market-intelligence",
         1, 6.0, 11, 0.4, size=12, color=CYAN, align=PP_ALIGN.CENTER)

# Save
os.makedirs("presentation", exist_ok=True)
prs.save("presentation/Bangkok_Airbnb_Market_Intelligence.pptx")
print("✅ Presentation saved: presentation/Bangkok_Airbnb_Market_Intelligence.pptx")
print(f"   Slides: {len(prs.slides)}")